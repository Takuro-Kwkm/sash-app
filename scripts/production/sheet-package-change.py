#!/usr/bin/env python3
"""Native workbook changes; the acquired Shared Harness owns the Work lifecycle."""
import argparse
import json
import sys
import subprocess
from functools import lru_cache
from pathlib import Path

STAGES = ('BASELINE_RESOLUTION', 'SOURCE_DELTA', 'IMPACT_SCOPE', 'CARRY_FORWARD',
          'CHANGE_APPLICATION', 'NATIVE_QA', 'QA_READY', 'REVIEW_PACKET')


def setup(harness_root):
    sys.path.insert(0, str(Path(harness_root).resolve()))
    global require, sha, read_json, write_json, Blocked, SHARED_ROOT
    SHARED_ROOT = str(Path(harness_root).resolve())
    from harness.core import require, sha, read_json, write_json, Blocked


def verified(ref):
    path = Path(ref['path'])
    require(path.is_file() and not path.is_symlink(), 'NATIVE_INPUT_REQUIRED', str(path), 'INPUT')
    data = path.read_bytes()
    require(sha(data) == ref['sha256'], 'NATIVE_INPUT_CHANGED', str(path), 'INPUT')
    return data


def json_adapter(c, profile):
    if profile.get('baseline_kind') != 'EXISTING_FORMAL_JSON': return None
    import importlib.util
    ref = profile['domain_adapter']
    path = Path(c['checkout']) / ref['path']
    verified({'path':str(path), 'sha256':ref['sha256']})
    spec = importlib.util.spec_from_file_location('json_change_domain', path)
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
    module.setup(SHARED_ROOT)
    return module


def workbook(ref):
    verified(ref)
    return _workbook(ref['path'], ref['sha256'])


@lru_cache(maxsize=4)
def _workbook(path, digest):
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        cells = {}
        for sheet in book:
            cells[sheet.title] = {}
            for row in sheet:
                for cell in row:
                    if cell.value is None:
                        continue
                    require(cell.data_type != 'e', 'NATIVE_FORMULA_ERROR', sheet.title + '!' + cell.coordinate, 'QA')
                    if cell.value is not None:
                        value = cell.value
                        if hasattr(value, 'isoformat'):
                            value = {'date': value.isoformat()}
                        cells[sheet.title][cell.coordinate] = value
        return cells
    finally:
        book.close()


def delta(before, after):
    changes = {}
    for sheet in sorted(before.keys() | after.keys()):
        old, new = before.get(sheet, {}), after.get(sheet, {})
        for cell in sorted(old.keys() | new.keys()):
            if old.get(cell) != new.get(cell):
                changes[sheet + '!' + cell] = {'before': old.get(cell), 'after': new.get(cell)}
    return changes


def baseline(c, profile):
    adapter = json_adapter(c, profile)
    if adapter:
        row = adapter.baseline(c, profile)
        manifest = json.loads(verified(c["baseline_live_manifest"]))
        return {**manifest, "package_version": manifest.get("formal_revision", manifest.get("package_version"))}, row
    from harness.production import native_registry, connector_document
    row = native_registry(json.loads(verified(c['native_registry'])), profile['manufacturer'], profile['registry_series'])
    manifest = json.loads(verified(c['baseline_manifest']))
    require(row['package_version'] == manifest['package_version'] == profile['formal_revision']
            and row['authoring_file_id'] == manifest['authoring_file_id']
            and row['runtime_manifest_id'] == c['baseline_manifest_id']
            and manifest['canonical_product_id'] == profile['canonical_product_id']
            and manifest['status'] == 'FORMAL_PASS' and manifest['formal_pass'] is True,
            'BASELINE_IDENTITY', 'Reacquire the exact native Current Formal', 'BASELINE_RESOLUTION')
    metadata = connector_document(json.loads(verified(c['baseline_authoring_metadata'])))
    require(metadata['id'] == row['authoring_file_id'] and metadata['parent_ids'] == [manifest['canonical_folder_id']]
            and metadata['mime_type'] == 'application/vnd.google-apps.spreadsheet',
            'AUTHORING_IDENTITY', 'Current Authoring location/type differs', 'BASELINE_RESOLUTION')
    documents = {}
    for item in manifest['runtime_files'] + [manifest['schema']]:
        data = verified(c['baseline_documents'][item['file_name']])
        require(sha(data) == item['sha256'] and len(data) == item['byte_size'], 'PACKAGE_HASH', item['file_name'], 'BASELINE_RESOLUTION')
        documents[item['file_name']] = json.loads(data)
    from jsonschema import Draft202012Validator
    schema = documents[manifest['schema']['file_name']]
    for item in manifest['runtime_files']:
        require(not list(Draft202012Validator(schema).iter_errors(documents[item['file_name']])),
                'BASELINE_SCHEMA', item['file_name'], 'BASELINE_RESOLUTION')
    require(list(workbook(c['baseline_authoring'])) == profile['native_sheets'],
            'NATIVE_SHEET_SCOPE', 'Current complete workbook required', 'BASELINE_RESOLUTION')
    return manifest, row


def review(c, profile):
    adapter = json_adapter(c, profile)
    if adapter:
        value = json.loads(verified(c["review"]))
        adapter.validate_change(json.loads(verified(c["baseline_master"])), json.loads(verified(c["candidate_master"])), value)
        require(value["product_id"] == c["product_id"] and value["scope_id"] == c["scope_id"], "SOURCE_BINDING", "Wrong review scope", "SOURCE_DELTA")
        sources = {r["sha256"] for r in c["sources"]}
        for ref in c["sources"]: verified(ref)
        require(all(r["source_sha256"] in sources for r in value["claims"]), "SOURCE_BINDING", "Uncaptured claim source", "SOURCE_DELTA")
        return value, value["delta"]
    value = json.loads(verified(c['review']))
    require(value['status'] == 'COMPLETE_REVIEWED' and value.get('review_ref')
            and value['product_id'] == c['product_id'] and value['scope_id'] == c['scope_id']
            and value['baseline_sha256'] == c['baseline_authoring']['sha256']
            and value['candidate_sha256'] == c['candidate_authoring']['sha256'],
            'SOURCE_REVIEW_REQUIRED', 'Exact complete source-bound review required', 'SOURCE_DELTA')
    require(set(value['sections']) == set(profile['review_sections'])
            and all(v == 'REVIEWED' for v in value['sections'].values()) and not value['unresolved_in_scope'],
            'SOURCE_REVIEW_REQUIRED', 'Review every business responsibility section', 'SOURCE_DELTA')
    require(value['source_hashes'] == {s['source_id']: s['sha256'] for s in c['sources']},
            'SOURCE_BINDING', 'Every source must be captured', 'SOURCE_DELTA')
    for source in c['sources']:
        verified(source)
    changes = delta(workbook(c['baseline_authoring']), workbook(c['candidate_authoring']))
    require(changes == value['cell_changes'], 'UNREVIEWED_CELL_CHANGE', 'Workbook differs from reviewed delta', 'SOURCE_DELTA')
    require(changes and value['changed_fields'] and value['evidence_ids'] and value['locators'],
            'EMPTY_CHANGE', 'Changed Work requires reviewed product changes', 'SOURCE_DELTA')
    require(set(profile['native_sheets']).issubset(workbook(c['candidate_authoring'])),
            'NATIVE_SHEET_REMOVAL', 'Retain existing Authoring sheets', 'CARRY_FORWARD')
    return value, changes


def candidate_parity(c):
    from openpyxl.utils import get_column_letter
    master=json.loads(verified(c['candidate_master']))
    cells=workbook(c['candidate_authoring'])
    checked=0
    for binding in c['sheet_bindings']:
        records=master
        for part in binding['records_pointer'].strip('/').split('/'):
            records=records[part]
        records=records[binding.get('record_offset',0):]
        if 'record_count' in binding: records=records[:binding['record_count']]
        sheet=cells[binding['sheet']]
        for row_index,record in enumerate(records,binding['start_row']):
            for column,key in enumerate(record,1):
                letter=get_column_letter(column)
                require(sheet.get(letter+str(binding['header_row']))==key, 'CANDIDATE_HEADER_PARITY', binding['sheet']+':'+key, 'NATIVE_QA')
                actual=sheet.get(letter+str(row_index))
                expected=record[key]
                if isinstance(expected,(dict,list)):
                    actual=json.loads(actual)
                require(actual==expected, 'CANDIDATE_SHEET_PARITY', binding['sheet']+'!'+letter+str(row_index), 'NATIVE_QA')
                checked+=1
    return checked


def stage(c, gate, work):
    from harness.production import COMPONENTS
    profile = json.loads(verified(c['profile']))
    is_json = profile.get("baseline_kind") == "EXISTING_FORMAL_JSON"
    require(c['product_id'] == profile['product_id'], 'NATIVE_PRODUCT', 'Wrong native product', gate)
    manifest, row = baseline(c, profile)
    reviewed, changes = review(c, profile)
    result = {'status': 'PASS', 'product_id': c['product_id'], 'scope_id': c['scope_id'],
              'formal_mutations': 0, 'runtime_ui_mutations': 0, 'source_adoptions': 0}
    previous = lambda name: read_json(Path(work) / (name + '.json'))
    if gate == 'BASELINE_RESOLUTION':
        result.update(formal_revision=manifest['package_version'], native_registry_entry=row,
                      canonical_product_id=manifest['canonical_product_id'], authoring_identity='LIVE_NATIVE_JSON' if is_json else 'LIVE_NATIVE_SHEETS_EXPORT')
    elif gate == 'SOURCE_DELTA':
        result.update(classification='MODIFIED', review_ref=reviewed['review_ref'], findings=reviewed['findings'],
                      source_role='NEW_CANDIDATE', changed_cells=len(changes))
    elif gate == 'IMPACT_SCOPE':
        require(set(reviewed['components']) == set(COMPONENTS) and not reviewed['unknown_impact'],
                'IMPACT_MAPPING_REQUIRED', 'Complete reverse consumer coverage required', gate)
        result.update(changed_fields=reviewed['changed_fields'], components=reviewed['components'])
    elif gate == 'CARRY_FORWARD':
        if is_json:
            from harness.core import package_members, inspect_package
            package = verified(c['baseline_package'])
            info = inspect_package(package); _, files = package_members(package)
            require(info['revision'] == profile['baseline_revision'] and files['PRODUCT_MASTER.json'] == verified(c['baseline_master'])
                    and files['CARRIED_CONTENT_QA.json'] == verified(c['carried_content_qa'])
                    and files['FORMAL_CONTENT_QA.json'] == verified(c['carried_formal_qa']),
                    'CARRY_FORWARD_QA_CHANGED', 'Exact adopted package/QA required', gate)
            result.update(existing_qa='BYTE_IDENTICAL', controlled_unresolved='BYTE_IDENTICAL', baseline_formal='PRESERVED')
        else:
            old, new = workbook(c['baseline_authoring']), workbook(c['candidate_authoring'])
            retained = sum(value == new[sheet].get(cell) for sheet, cells in old.items() for cell, value in cells.items())
            result.update(existing_sheets=len(old), unchanged_cells=retained,
                          evidence_ids='PRESERVED', accepted_decisions='PRESERVED', baseline_formal='PRESERVED')
    elif gate == 'CHANGE_APPLICATION':
        destination = Path(work) / 'candidate'
        destination.mkdir(exist_ok=True)
        for name, ref in c['candidate_files'].items():
            require(Path(name).name == name, 'PACKAGE_PATH', 'Sibling files only', gate)
            data = verified(ref)
            path = destination / name
            require(not path.exists() or path.read_bytes() == data, 'IMMUTABLE_OUTPUT_CHANGED', name, gate)
            path.write_bytes(data)
        result.update(candidate_revision=c['candidate_revision'], files={k:v['sha256'] for k,v in c['candidate_files'].items()},
                      adoption='NOT_ADOPTED', runtime_scope='NOT_INCLUDED_IN_PRODUCT_MASTER_FORMAL')
    elif gate == 'NATIVE_QA':
        qa_path = Path(work) / 'project-qa.json'
        verified(c['qa_worker'])
        verified(c['qa_input'])
        if is_json:
            qa_input = read_json(c['qa_input']['path'])
            require(c['qa_input']['sha256'] == reviewed['qa_input_sha256'] and qa_input['candidate'] == c['candidate_master']
                    and qa_input['assertions'] and qa_input['business_cases'], 'NATIVE_QA_BINDING', 'Exact candidate and substantive QA required', gate)
            parity = 0
        else:
            parity=candidate_parity(c)
        subprocess.run([c['node_executable'], c['qa_worker']['path'], '--input', c['qa_input']['path'], '--out', str(qa_path)], check=True)
        qa = read_json(qa_path)
        require(qa['status'] == 'PASS' and qa['product_master_sha256' if is_json else 'candidate_sha256'] == c['candidate_master' if is_json else 'candidate_authoring']['sha256']
                and qa['failed'] == 0 and qa['executed'] is True and qa['tests'],
                'NATIVE_QA_REQUIRED', 'Same-candidate executed business QA required', gate)
        result.update(qa=qa, candidate_parity_cells=parity)
    elif gate == 'QA_READY':
        qa = previous('NATIVE_QA')['qa']
        result.update(QA_READY=True, candidate_revision=c['candidate_revision'], business_scope='SALES_ESTIMATE_REQUEST',
                      passed=qa['passed'], controlled_unresolved=reviewed['controlled_unresolved'],
                      existing_formal='PRESERVED', formal='NOT_ADOPTED')
    elif gate == 'REVIEW_PACKET':
        result.update(human_required=False, adoption_requested=False, review_ref=reviewed['review_ref'],
                      candidate=previous('CHANGE_APPLICATION'), qa=previous('QA_READY'),
                      formal_policy='A future Formal request needs a fixed packet and an actual same-payload external Human Decision')
    else:
        raise ValueError('Unknown native stage: ' + gate)
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for key in ('harness-root', 'contract', 'stage', 'work', 'out'):
        parser.add_argument('--' + key, required=True)
    args = parser.parse_args()
    setup(args.harness_root)
    try:
        value = stage(read_json(args.contract), args.stage, args.work)
    except (Blocked, OSError, ValueError, KeyError, ImportError, subprocess.CalledProcessError) as error:
        value = {'status': 'FAIL', 'blocking_reason': error.data if isinstance(error, Blocked) else {'code':'NATIVE_CHANGE_FAILURE','reason':str(error)}}
    write_json(args.out, value)
