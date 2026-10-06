#!/usr/bin/env python3
"""ThermosL native workers; shared Workflow owns execution and Resume."""
import argparse
import json
import subprocess
import sys
from pathlib import Path


def execute(stage, contract, work, harness_root):
    sys.path.insert(0, str(Path(harness_root).resolve()))
    from harness.core import read_json, sha, require, write_json, fingerprint
    from harness.changed_path import checked_bytes
    from harness.production import short_route
    c, work = read_json(contract), Path(work)
    result = {'status': 'PASS', 'stage': stage, 'product_id': c['product_id'], 'scope': c['scope']}
    if stage == 'ROUTING':
        result['route'] = short_route(c['instruction'], c['routing'])
    elif stage == 'CURRENT_FORMAL_RESOLUTION':
        checked_bytes(c['master_before']); checked_bytes(c['runtime_before'])
        require(c['registry_entry']['package_version'] == 'v0.7-R2', 'FORMAL_BASELINE_CHANGED', 'Acquire current native Formal', stage)
        result['formal_version'] = 'v0.7-R2'
    elif stage == 'SOURCE_DELTA':
        # Actual complete PDFs, not a generated manufacturer fixture.
        import fitz
        for key in ('source_before', 'source_after'): checked_bytes(c[key])
        before = fitz.open(c['source_before']['path']); after = fitz.open(c['source_after']['path'])
        label = '代替進入口用幅木仕様'
        old = [p.get_text() for p in before]; new = [p.get_text() for p in after]
        require(not any(label in p for p in old), 'DELTA_BASELINE', 'New option already present; inspect lineage', stage)
        require(all(label in new[p-1] and '標準タイプのみ' in new[p-1] for p in (57, 61)), 'OFFICIAL_DELTA', 'Exact official claim missing', stage)
        result.update(classifications=['ADDED', 'CONSTRAINT_CHANGE'], changed_values=['OP-SL-EMERGENCY-SKIRT'],
                      accepted_pdf_pages=[57, 61], accepted_printed_pages=[55, 59],
                      differing_text_pages=[i+1 for i in range(len(old)) if old[i] != new[i]],
                      catalogue_review_scope='ONE_OPTION_ONLY; other catalogue changes are not adopted',
                      unknown_outside_scope=c['outside_scope_review'])
    elif stage == 'IMPACT_SCOPE':
        result.update(components=c['validation']['components'], classification='RUNTIME_IMPACT', field_id='sash-app:SER-LIX-SAMOSL:options')
    elif stage in ('CARRY_FORWARD', 'VALIDATION'):
        import openpyxl
        old = openpyxl.load_workbook(c['master_before']['path']); new = openpyxl.load_workbook(c['master_candidate']['path'])
        require(old.sheetnames == new.sheetnames, 'VALIDATOR_FAILURE', 'Native sheet identity differs', stage)
        mismatches = []
        for name in old.sheetnames:
            for row in old[name].iter_rows():
                for cell in row:
                    if new[name][cell.coordinate].value != cell.value: mismatches.append((name, cell.coordinate))
        require(not mismatches, 'VALIDATOR_FAILURE', 'Unexpected existing-cell mutation: '+str(mismatches[:5]), stage)
        require(new['10_その他OP']['A201'].value == 'OP-SL-EMERGENCY-SKIRT'
                and new['10_その他OP']['D201'].value == 'SP-SL-SHUT-M-STD', 'VALIDATOR_FAILURE', 'Scoped option identity differs', stage)
        checked_bytes(c['master_candidate']); checked_bytes(c['runtime_candidate'])
        result.update(existing_cells_preserved=c['validation']['preserved_cells'], sheets_preserved=41, prior_fields_preserved=31, changed_fields=1)
    elif stage == 'PRODUCT_MASTER_CHANGE':
        from harness.artifacts import save_verified
        result['artifacts'] = [save_verified(Path(c[key]['path']), work / name) for key, name in
                              [('master_candidate', 'selected-master.xlsx'), ('runtime_candidate', 'selected-runtime.json')]]
    elif stage == 'RUNTIME_UI_VALIDATION':
        proc = subprocess.run(['node', '--test', c['runtime_test']], capture_output=True, text=True, timeout=120)
        result['actual_runtime_ui_test_output'] = proc.stdout + proc.stderr
        require(proc.returncode == 0, 'VALIDATOR_FAILURE', result['actual_runtime_ui_test_output'][-1500:], stage)
    elif stage == 'QA_READY':
        require(c['validation']['existing_cell_differences'] == [], 'VALIDATOR_FAILURE', 'Carry-forward is not qualified', stage)
        result.update(state='QA_READY_SCOPED', estimate_confirmation_required=True, promotion_eligible=True,
                      formal_adoption='HUMAN_DECISION_REQUIRED', all_catalogue_adoption=False)
    elif stage == 'FORMAL_REVIEW':
        result.update(human_required=True, human_decision='NOT_ACCEPTED', candidate_artifacts=c['review_artifacts'],
                      proposed_formal_transition=c['formal_transition'], downstream_plan=c['downstream_plan'],
                      required_native_save=c['native_save_requests'], architecture_deviation=0)
        result['payload_sha256'] = fingerprint(result)
    elif stage == 'FORMAL_ADOPTION':
        from harness.changed_path import save_formal
        require(c.get('formal_plan'), 'FORMAL_PLAN_REQUIRED', 'Complete the approved exact native Formal transaction', stage)
        result = save_formal(c['formal_root'], read_json(c['formal_plan']),
            read_json(work/'review-packets/HUMAN_ADOPTION.json'), read_json(c['decision']), read_json(c['native_receipt']))
    elif stage == 'DOWNSTREAM_UPDATE':
        from harness.changed_path import apply_downstream
        result = apply_downstream(c['checkout'], work, c['downstream_plan'], read_json(work/'FORMAL_ADOPTION.json'))
    else:
        raise ValueError('Unknown native stage: '+stage)
    return result


if __name__ == '__main__':
    p = argparse.ArgumentParser(); p.add_argument('--harness-root', required=True); p.add_argument('--contract', required=True)
    p.add_argument('--stage', required=True); p.add_argument('--work', required=True); p.add_argument('--out', required=True)
    a = p.parse_args()
    sys.path.insert(0, str(Path(a.harness_root).resolve()))
    from harness.core import write_json, Blocked
    try:
        value = execute(a.stage, a.contract, a.work, a.harness_root)
    except (Blocked, OSError, ValueError, KeyError) as e:
        value = {'status': 'FAIL', 'blocking_reason': e.data if isinstance(e, Blocked) else {'reason': str(e)}}
    write_json(a.out, value)
