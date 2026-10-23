"""Synthetic negative checks for the native source-format boundary."""
import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook

definition = importlib.util.spec_from_file_location('xlsx_recovery', Path(__file__).with_name('working-master-xlsx-recovery.py'))
adapter = importlib.util.module_from_spec(definition)
definition.loader.exec_module(adapter)
adapter.setup(sys.argv.pop(1))


class BoundarySafety(unittest.TestCase):
    def workbook(self, header, row):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        path = Path(temp.name) / 'source.xlsx'
        book = Workbook()
        sheet = book.active
        for r in [['Original title'], ['Original revision'], [], header, row]:
            sheet.append(r)
        book.save(path)
        book.close()
        return path

    def test_header_offset_keeps_preamble_and_null(self):
        path = self.workbook(['id', 'unknown', 'status'], ['A', None, 'HOLD'])
        result = adapter.tables(path, 4)['Sheet']
        self.assertEqual(result['rows'], [{'id': 'A', 'unknown': None, 'status': 'HOLD'}])
        self.assertEqual(result['preamble'][0][0], 'Original title')

    def test_wrong_header_is_not_silently_repaired(self):
        with self.assertRaises(adapter.Blocked):
            adapter.tables(self.workbook(['id', 'value'], ['A', 2]), 1)

    def test_duplicate_header_is_rejected(self):
        with self.assertRaises(adapter.Blocked):
            adapter.tables(self.workbook(['id', 'id'], ['A', 'B']), 4)

    def test_undeclared_column_is_not_dropped(self):
        with self.assertRaises(adapter.Blocked):
            adapter.tables(self.workbook(['id'], ['A', 'not declared']), 4)


if __name__ == '__main__':
    unittest.main()
