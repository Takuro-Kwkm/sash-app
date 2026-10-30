#!/usr/bin/env python3
"""Lossless native authoring intake. No product rules or lifecycle engine."""
import argparse, csv, io, json, sys, hashlib
from pathlib import Path

def inspect_workbook(path, profile):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        tables = {}
        for sheet in book:
            sheet.reset_dimensions()
            cells = {cell.coordinate:{'value':cell.value, 'type':cell.data_type}
                     for row in sheet.iter_rows() for cell in row if cell.value is not None}
            tables[sheet.title] = {'cells':cells}
        if set(tables) != set(profile['required_sheets']): raise ValueError('Complete native workbook inventory differs')
        return {'schema_version':'NATIVE_AUTHORING_CELL_SNAPSHOT_V1','product_id':profile['product_id'],
            'manufacturer':profile['manufacturer'],'series':profile['registry_series'],
            'source_revision':profile['baseline_revision'],'source_sha256':hashlib.sha256(Path(path).read_bytes()).hexdigest(),
            'tables':tables,'qa_ready':False,'formal_pass':False,'automatic_orderability':False,
            'historical_status_policy':'Cell contents and legacy PASS labels are preserved history, not current QA'}
    finally: book.close()

def verify_inputs(contract, profile, shared):
    sys.path.insert(0,str(Path(shared).resolve()))
    from harness.core import require, read_json, sha
    from harness.resolvers import capture
    for ref in contract['inputs'].values():capture(Path('/'),ref)
    raw=read_json(contract['inputs']['registry']['path'])
    require(not raw.get('isError'),'NATIVE_REGISTRY_ACQUISITION','Registry failure is not absence','INTAKE')
    data=raw.get('structuredContent',raw)
    require(data['id']==profile['registry_id'],'NATIVE_REGISTRY_IDENTITY','Wrong Registry','INTAKE')
    rows=list(csv.DictReader(io.StringIO(data['content'])))
    matches=[r for r in rows if r.get('manufacturer')==profile['manufacturer'] and r.get('series')==profile['registry_series']]
    require(len(matches)==1 and matches[0]['master_status'] in profile['baseline_states']
        and matches[0]['authoring_file_id']==profile['authoring_file_id']
        and matches[0]['package_version']==profile['baseline_revision'],
        'NATIVE_EXISTING_AUTHORING_REQUIRED','Preserve the real existing package-blocked row','INTAKE')
    row=matches[0]
    require(all(row[k]==profile[k] for k in ('canonical_folder_id','working_folder_id','old_folder_id')),
        'NATIVE_FOLDER_DRIFT','Actual registered folders required','INTAKE')
    for key in ('authoring_metadata','authoring_readback_metadata'):
        raw=read_json(contract['inputs'][key]['path']);meta=raw.get('structuredContent',raw)
        require(not raw.get('isError') and meta.get('id')==profile['authoring_file_id']
            and meta.get('parent_ids')==[profile['canonical_folder_id']],
            'NATIVE_AUTHORING_LOCATION','Original canonical authoring identity differs','INTAKE')
    original=capture(Path('/'),contract['inputs']['authoring'])
    require(original==capture(Path('/'),contract['inputs']['authoring_readback'])
        and sha(original)==profile['authoring_sha256'],'NATIVE_AUTHORING_BYTE_PARITY','Exact acquired original required','INTAKE')
    restored=inspect_workbook(contract['inputs']['authoring']['path'],profile)
    require(read_json(contract['inputs']['original']['path'])==restored,
        'NATIVE_AUTHORING_CELL_PARITY','Every cell/formula must remain in the DRAFT snapshot','INTAKE')
    return restored

def execute(shared,contract_path,work,out):
    sys.path.insert(0,str(Path(shared).resolve()))
    from harness.core import read_json,write_json,require
    from harness.resolvers import capture
    c=read_json(contract_path);p=json.loads(capture(Path('/'),c['profile']))
    work,out=Path(work).resolve(),Path(out).resolve()
    require(out.is_relative_to(work) and not out.is_symlink(),'OUTPUT_PATH','Regular Work-owned output required','INTAKE')
    snapshot=verify_inputs(c,p,shared)
    result={'status':'PASS','scope':'LOSSLESS_INPUT_RECOVERY_ONLY','sheets':len(snapshot['tables']),
        'cells':sum(len(x['cells']) for x in snapshot['tables'].values()),'qa_ready':False,'formal':False}
    write_json(out,result);return result

if __name__=='__main__':
    a=argparse.ArgumentParser()
    for k in ('harness-root','contract','work','out'):a.add_argument('--'+k,required=True)
    v=a.parse_args()
    try: result=execute(v.harness_root,v.contract,v.work,v.out)
    except Exception as error:
        result={'status':'FAIL','blocking_reason':getattr(error,'data',{'reason':str(error)})}
        Path(v.out).write_text(json.dumps(result,ensure_ascii=False,indent=2))
    raise SystemExit(0 if result['status']=='PASS' else 2)
