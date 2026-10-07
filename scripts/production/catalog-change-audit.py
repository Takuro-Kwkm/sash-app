#!/usr/bin/env python3
"""Native, read-only catalogue comparison; the shared v2 engine owns the DAG.

Product identity and source/review scope come from immutable native data. This
worker never applies product facts, adopts sources or issues a Human Decision.
"""
import argparse
import json
import re
import sys
from pathlib import Path

STAGES = ('BASELINE_RESOLUTION', 'SOURCE_DELTA', 'IMPACT_SCOPE', 'CARRY_FORWARD',
          'NATIVE_QA', 'AUDIT_REVIEW', 'COMPLETION')
SECTIONS = ('identity_and_effective_date', 'product_and_specification',
            'dimensions_and_availability', 'glass_and_performance',
            'colors_and_screens', 'options_and_dependencies',
            'construction_and_required_parts', 'codes_and_prices',
            'delivery_warranty_and_notes')


def setup(harness_root):
    sys.path.insert(0, str(Path(harness_root).resolve()))
    global require, read_json, sha, write_json, Blocked
    from harness.core import require, read_json, sha, write_json, Blocked


def verified(ref):
    path = Path(ref['path'])
    require(path.is_file() and not path.is_symlink(), 'NATIVE_REF_MISSING', str(path), 'ACQUISITION')
    data = path.read_bytes()
    require(sha(data) == ref['sha256'], 'NATIVE_REF_CHANGED', str(path), 'ACQUISITION')
    return data


def compare_pages(before, after):
    """Keep extraction separate from semantic review; price masking is triage."""
    import pypdfium2 as pdf
    docs = [pdf.PdfDocument(verified(ref)) for ref in (before, after)]
    try:
        texts = []
        for doc in docs:
            pages = []
            for page in doc:
                text = page.get_textpage()
                try:
                    pages.append(text.get_text_range())
                finally:
                    text.close(); page.close()
            texts.append(pages)
        old, new = texts
        pages = []
        for i in range(max(len(old), len(new))):
            a = old[i] if i < len(old) else None
            b = new[i] if i < len(new) else None
            masked = lambda t: re.sub(r'[¥￥]\s*[\d,]+', '<PRICE>', t) if t is not None else None
            pages.append({'pdf_page': i + 1, 'before_text_sha256': sha(a.encode()) if a is not None else None,
                          'after_text_sha256': sha(b.encode()) if b is not None else None,
                          'text_changed': a != b, 'nonprice_text_changed': masked(a) != masked(b),
                          'extraction_warning': any('\ufffd' in t for t in (a or '', b or ''))})
        return {'before_page_count': len(old), 'after_page_count': len(new), 'pages': pages,
                'text_changed_pages': [p['pdf_page'] for p in pages if p['text_changed']],
                'nonprice_text_changed_pages': [p['pdf_page'] for p in pages if p['nonprice_text_changed']],
                'semantic_inference': 'NONE; rendered tables and notes require bound review'}
    finally:
        for doc in docs: doc.close()


def admitted_review(c, comparison):
    review = json.loads(verified(c['review']))
    require(review.get('status') == 'COMPLETE_REVIEWED' and review.get('review_ref'),
            'SOURCE_REVIEW_REQUIRED', 'A complete native semantic review is required', 'SOURCE_DELTA')
    require(review['product_id'] == c['product_id'] and review['scope_id'] == c['scope_id']
            and review['before_sha256'] == c['source_before']['sha256']
            and review['after_sha256'] == c['source_after']['sha256'],
            'REVIEW_IDENTITY', 'Review is for a different product/source/scope', 'SOURCE_DELTA')
    require(set(review['sections']) == set(SECTIONS) and all(v == 'REVIEWED' for v in review['sections'].values()),
            'REVIEW_SCOPE', 'Complete declared catalogue responsibility scope required', 'SOURCE_DELTA')
    require(review['compared_pdf_pages'] == list(range(1, max(comparison['before_page_count'], comparison['after_page_count']) + 1)),
            'REVIEW_PAGE_SCOPE', 'Every PDF page must be compared', 'SOURCE_DELTA')
    require(sorted(review['reviewed_nonprice_pdf_pages']) == comparison['nonprice_text_changed_pages'],
            'REVIEW_DELTA_DRIFT', 'Unreviewed non-price text change', 'SOURCE_DELTA')
    visual = json.loads(verified(c['visual_comparison']))
    require(visual['before_sha256'] == review['before_sha256'] and visual['after_sha256'] == review['after_sha256']
            and [p['pdf_page'] for p in visual['pages']] == list(range(1, len(comparison['pages']) + 1)),
            'VISUAL_SCOPE', 'Complete source-bound rendered comparison required', 'SOURCE_DELTA')
    require(sorted(review['reviewed_visual_pdf_pages']) == [p['pdf_page'] for p in visual['pages'] if p['nonprice_visual_changed']],
            'VISUAL_REVIEW_REQUIRED', 'Every unmasked rendered difference must be reviewed', 'SOURCE_DELTA')
    require(not review.get('unresolved_in_scope'), 'SOURCE_REVIEW_REQUIRED', 'Catalogue audit is not fully classified', 'SOURCE_DELTA')
    return review


def stage(c, gate, work):
    from harness.production import native_registry, COMPONENTS
    work = Path(work)
    previous = lambda name: read_json(work / (name + '.json'))
    result = {'status': 'PASS', 'product_id': c['product_id'], 'scope_id': c['scope_id'],
              'formal_mutations': 0, 'source_adoptions': 0, 'runtime_mutations': 0}
    if gate == 'BASELINE_RESOLUTION':
        profile = json.loads(verified(c['profile']))
        require(profile['product_id'] == c['product_id'] and profile['mutation_policy'] == 'NONE',
                'NATIVE_AUDIT_SCOPE', 'Profile product/mutation boundary differs', gate)
        row = native_registry(json.loads(verified(c['native_registry'])), profile['manufacturer'], profile['registry_series'])
        manifest = json.loads(verified(c['runtime_manifest']))
        require(row['package_version'] == manifest['package_version'] == profile['formal_revision']
                and row['authoring_file_id'] == manifest['authoring_file_id'] == c['authoring_file_id']
                and row['runtime_manifest_id'] == c['runtime_manifest_file_id']
                and manifest['formal_status'] == 'FORMAL_PASS',
                'BASELINE_IDENTITY', 'Current Registry, live manifest and native profile differ', gate)
        require(sha(verified(c['authoring'])) == manifest['authoring_sha256'],
                'AUTHORING_FORMAL_DRIFT', 'Native workbook differs from live Formal manifest', gate)
        from harness.production import connector_document
        for key in ('source_before', 'source_after'):
            metadata = connector_document(json.loads(verified(c[key+'_metadata'])))
            require(metadata['id'] == profile[key]['file_id'] and metadata['mime_type'] == 'application/pdf'
                    and int(metadata['size']) == len(verified(c[key])),
                    'SOURCE_METADATA', 'Full source identity/MIME/size differs', gate)
        locator = json.loads(verified(c['official_locator']))
        require(locator['observation_type'] == 'RENDERED_OFFICIAL_CATALOGUE' and locator['code'] == profile['source_after']['code']
                and locator['issue'] == profile['source_after']['issue'] and locator['price_edition'] == profile['source_after']['price_edition']
                and locator['url'] == profile['source_discovery']['official_catalogue_url'],
                'LATEST_SOURCE_DISCOVERY_REQUIRED', 'Reacquire official Current locator; source filenames do not prove latest', gate)
        for ref in c['preserved_refs']: verified(ref)
        result.update(formal_revision=row['package_version'], authoring_sha256=c['authoring']['sha256'],
                      runtime_manifest_sha256=c['runtime_manifest']['sha256'], native_registry_entry=row,
                      authoring_acquisition='LOCAL_MIRROR_BYTES_EQUAL_LIVE_MANIFEST_SHA256')
    elif gate == 'SOURCE_DELTA':
        comparison = compare_pages(c['source_before'], c['source_after'])
        require(comparison['after_page_count'] == json.loads(verified(c['official_locator']))['page_count'],
                'SOURCE_PAGE_SCOPE', 'Full PDF differs from live official catalogue page count', gate)
        review = admitted_review(c, comparison)
        result.update(comparison=comparison, classification='MODIFIED' if c['source_before']['sha256'] != c['source_after']['sha256'] else 'UNCHANGED',
                      findings=review['findings'], review_ref=review['review_ref'],
                      source_role='NEW_CANDIDATE; existing adopted Source unchanged',
                      scope_limits=review['scope_limits'])
    elif gate == 'IMPACT_SCOPE':
        delta = previous('SOURCE_DELTA')
        require(all(f['classification'] in ('ADDED', 'REMOVED', 'MODIFIED', 'CONSTRAINT_CHANGE', 'PRESENTATION_ONLY', 'SOURCE_CHANGED')
                    and f['affected_sheets'] and f['pdf_pages'] for f in delta['findings']),
                'IMPACT_MAPPING_REQUIRED', 'Every finding needs native sheet and source mapping', gate)
        result.update(affected_sheets=sorted({s for f in delta['findings'] for s in f['affected_sheets']}),
                      components={name: 'REVIEW_CANDIDATE_ONLY' if name in ('Product Master', 'Evidence', 'Dependency', 'Constraints', 'Selection Contract', 'Canonical Mapping') else 'NO_WRITE_IN_AUDIT_SCOPE' for name in COMPONENTS},
                      application_state='NOT_APPLIED; existing fields/Evidence/accepted decisions retained')
    elif gate == 'CARRY_FORWARD':
        for ref in c['preserved_refs']: verified(ref)
        result.update(preserved_files=len(c['preserved_refs']), formal_revision=previous('BASELINE_RESOLUTION')['formal_revision'],
                      controlled_unresolved='PRESERVED', evidence_ids='PRESERVED', accepted_decisions='PRESERVED')
    elif gate == 'NATIVE_QA':
        import openpyxl
        w = openpyxl.load_workbook(Path(c['authoring']['path']), read_only=True, data_only=False)
        profile = json.loads(verified(c['profile']))
        try:
            require(w.sheetnames == profile['native_sheets'], 'NATIVE_SHEET_SCOPE', 'Native sheet identity differs', gate)
            errors = [(s.title, cell.coordinate) for s in w for row in s for cell in row if cell.data_type == 'e']
            require(not errors, 'NATIVE_FORMULA_ERROR', str(errors[:5]), gate)
        finally: w.close()
        review = admitted_review(c, previous('SOURCE_DELTA')['comparison'])
        require(all(set(f['affected_sheets']).issubset(profile['native_sheets']) for f in review['findings']),
                'IMPACT_MAPPING_REQUIRED', 'Unknown native sheet', gate)
        for ref in c['preserved_refs']: verified(ref)
        result.update(native_sheets=len(profile['native_sheets']), formula_errors=0,
                      audit_validation='PASS', product_qa='NOT_REISSUED', formal_promotion='NOT_REQUESTED')
    elif gate == 'AUDIT_REVIEW':
        result.update(state='SOURCE_AUDIT_REVIEWED', delta=previous('SOURCE_DELTA'), impact=previous('IMPACT_SCOPE'),
                      next_action='Create a separately scoped Working candidate if application is requested; consume an actual Human Decision before Formal adoption')
    elif gate == 'COMPLETION':
        for ref in c['preserved_refs']: verified(ref)
        require(previous('NATIVE_QA')['audit_validation'] == 'PASS', 'AUDIT_QA_REQUIRED', 'Audit validation did not pass', gate)
        result.update(state='SOURCE_AUDIT_VERIFIED', findings=previous('SOURCE_DELTA')['findings'],
                      scope_limits=previous('SOURCE_DELTA')['scope_limits'], target_gate='SOURCE_AUDIT_VERIFIED',
                      existing_formal='PRESERVED', cloud_sync='NOT_VERIFIED')
    else: raise ValueError('Unknown audit stage: ' + gate)
    return result


if __name__ == '__main__':
    p = argparse.ArgumentParser()
    for key in ('harness-root', 'contract', 'stage', 'work', 'out'): p.add_argument('--' + key, required=True)
    a = p.parse_args(); setup(a.harness_root)
    try: value = stage(read_json(a.contract), a.stage, a.work)
    except (Blocked, OSError, ValueError, KeyError, ImportError) as e:
        value = {'status': 'FAIL', 'blocking_reason': e.data if isinstance(e, Blocked) else {'code': 'NATIVE_AUDIT_FAILURE', 'reason': str(e)}}
    write_json(a.out, value)
