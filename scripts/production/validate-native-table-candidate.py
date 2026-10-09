#!/usr/bin/env python3
"""Native Working-only parity/storage eligibility; never manufactures a QA_READY."""
import argparse
import hashlib
import json
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter


def check(c):
    root=Path(c['checkout']);folder=root/c['candidate_folder'];manifest=json.loads((folder/'candidate-manifest.json').read_text())
    assert manifest['schema']=='NATIVE_QUARANTINED_CORRECTION_PACKET_V1'
    assert manifest['qa_ready'] is False and manifest['formal_adopted'] is False and manifest['automatic_orderability'] is False
    for ref in manifest['members']:
        b=(folder/ref['path']).read_bytes()
        assert len(b)==ref['bytes'] and hashlib.sha256(b).hexdigest()==ref['sha256'],ref['path']
    for component in manifest['carry_forward']:
        for name in ('authoring','runtime'):
            ref=component[name];assert hashlib.sha256((root/ref['path']).read_bytes()).hexdigest()==ref['sha256']
    plan=json.loads((folder/'patch-plan.json').read_text());baseline=json.loads((root/c['baseline_runtime']).read_text());runtime=json.loads((folder/'runtime-candidate.json').read_text())
    assert hashlib.sha256((root/c['baseline_runtime']).read_bytes()).hexdigest()==plan['base_runtime_sha256']
    allowed={(p['sheet'],p['address']):p for p in plan['patches']}
    book=openpyxl.load_workbook(folder/'authoring-candidate.xlsx',data_only=True)
    changes=[]
    try:
        assert set(runtime['tables'])==set(baseline['tables'])==set(book.sheetnames)
        for name,table in runtime['tables'].items():
            before=baseline['tables'][name]['values']
            for row,values in enumerate(table['values'],1):
                for col,value in enumerate(values,1):
                    assert value==book[name].cell(row,col).value
                    old=before[row-1][col-1] if row<=len(before) and col<=len(before[0]) else None
                    if old!=value:
                        cell=get_column_letter(col)+str(row);patch=allowed[(name,cell)]
                        assert (old,value)==(patch['before'],patch['after'])
                        changes.append([name,cell])
        assert len(changes)==len(plan['patches'])
    finally:book.close()
    return {'status':'PASS','state':'WORKING_CANDIDATE_STORAGE_ELIGIBLE','patch_cells':len(changes),
            'carry_forward_components':len(manifest['carry_forward']),'qa_ready':False,'formal_adopted':False,
            'product_qa':'NOT_REISSUED','external_operations':0}


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--contract',required=True);p.add_argument('--stage',required=True,choices=['validate','source-gate']);p.add_argument('--out',required=True);a=p.parse_args()
    try:
        c=json.loads(Path(a.contract).read_text());result=check(c)
        if a.stage=='source-gate':
            result={'status':'FAIL','blocking_reason':{'code':'OFFICIAL_SOURCE_REVIEW_REQUIRED',
                'reason':'Complete all-responsibility official Source review and legacy reference closure in a new immutable successor; existing audit/storage PASS is not product QA.'},
                'completed_candidate_storage_preserved':True,'qa_ready':False,'external_operations':0}
    except (AssertionError,OSError,ValueError,KeyError,TypeError) as error:
        result={'status':'FAIL','blocking_reason':{'code':'NATIVE_CANDIDATE_REJECTED','reason':str(error)},'external_operations':0}
    Path(a.out).write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
