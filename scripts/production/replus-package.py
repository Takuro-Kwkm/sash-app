#!/usr/bin/env python3
"""Existing multi-Authoring native audit. Shared v2 owns DAG/save/resume.

The data profile owns identities and sheet topology. No manufacturer rule is
inferred from prose, old QA labels, or a successful package hash check.
"""
import argparse
import hashlib
import json
import re
from pathlib import Path

STAGES = ('BASELINE_RESOLUTION', 'NATIVE_TABLE_PARITY', 'EVIDENCE_INVENTORY',
          'REFERENCE_CLOSURE', 'REVIEW_PACKET', 'COMPLETION')


class Rejected(ValueError):
    def __init__(self, code, reason):
        self.code = code
        super().__init__(reason)


def require(condition, code, reason):
    if not condition:
        raise Rejected(code, reason)


def sha(data):
    return hashlib.sha256(data).hexdigest()


def verified(ref, root):
    name = Path(ref['path'])
    path = name if name.is_absolute() else Path(root) / name
    require(path.is_file() and not path.is_symlink(), 'NATIVE_REF_MISSING', str(path))
    data = path.read_bytes()
    require(sha(data) == ref['sha256'], 'NATIVE_REF_CHANGED', str(path))
    return data


def document(ref, root):
    return json.loads(verified(ref, root))


def load_profile(checkout, profile_ref):
    profile = document(profile_ref, checkout)
    require(profile['schema'] == 'NATIVE_EXISTING_MULTI_AUTHORING_PROFILE_V1'
            and profile['mutation_policy'] == 'NONE', 'NATIVE_PROFILE_SCOPE', 'Read-only native profile required')
    ids = [c['component_id'] for c in profile['components']]
    require(len(ids) == 7 and len(set(ids)) == 7, 'NATIVE_COMPONENT_SCOPE', 'Seven unique captured components required')
    require(profile['admitted_target_gate'] == 'REPLUS_PACKAGE_AUDITED', 'NATIVE_TARGET_SCOPE', 'Unsupported audit target')
    return profile


def baseline(checkout, profile, registry_csv):
    import csv
    import io
    rows = list(csv.DictReader(io.StringIO(registry_csv)))
    matches = [r for r in rows if r['manufacturer'] == profile['manufacturer'] and r['series'] == profile['registry_series']]
    require(len(matches) == 1, 'NATIVE_REGISTRY_AMBIGUOUS', 'Acquire one exact Current Formal row')
    row = matches[0]
    manifest = document(profile['manifests']['runtime-manifest.json'], checkout)
    index = document(profile['manifests']['runtime-index.json'], checkout)
    require(row['master_status'] == manifest['master_status'] == 'FORMAL_PASS'
            and row['package_version'] == manifest['package_version'] == index['package_version'] == profile['formal_package_version'],
            'NATIVE_FORMAL_IDENTITY', 'Registry/package version or state differs')
    require(index['component_count'] == 7 and len(index['components']) == 7
            and len(manifest['runtime_files']) == 8 and len(manifest['authoring_files']) == 7,
            'NATIVE_COMPONENT_SCOPE', 'Manifest and Index must include all seven')
    ids = [c['component_id'] for c in profile['components']]
    require(set(ids) == {c['component_id'] for c in index['components']}, 'NATIVE_COMPONENT_SCOPE', 'Index has another component set')
    require(row['authoring_file_id'].split(';') == [c['authoring_file_id'] for c in profile['components']],
            'NATIVE_FORMAL_IDENTITY', 'Current authoring File IDs differ')
    require(row['runtime_manifest_id'] == manifest['runtime_manifest_file_id']
            and row['documentation_file_id'] == manifest['documentation_file']['file_id']
            and row['canonical_folder_id'] == manifest['canonical_folder_id'],
            'NATIVE_FORMAL_IDENTITY', 'Current artifact IDs/placement differ')
    for component in profile['components']:
        cid = component['component_id']
        a = [x for x in manifest['authoring_files'] if x['component_id'] == cid]
        r = [x for x in manifest['runtime_files'] if x['component_id'] == cid]
        i = [x for x in index['components'] if x['component_id'] == cid]
        require(len(a) == len(r) == len(i) == 1, 'NATIVE_COMPONENT_SCOPE', cid + ': duplicated or missing identity')
        payload = document(component['runtime'], checkout)
        verified(component['authoring'], checkout)
        require(a[0]['file_id'] == i[0]['authoring_file_id'] == payload['authoring_source']['file_id'] == component['authoring_file_id']
                and r[0]['file_id'] == i[0]['runtime_file_id'] == component['runtime_file_id']
                and r[0]['sha256'] == i[0]['runtime_sha256'] == component['runtime']['sha256']
                and a[0]['component_version'] == i[0]['component_version'] == payload['component_version'] == component['component_version']
                and payload['component_id'] == cid and payload['package_version'] == profile['formal_package_version']
                and payload['runtime_contract'] == index['runtime_contract'] == 'replus_table_driven_v1',
                'NATIVE_COMPONENT_IDENTITY', cid + ': Authoring/Runtime/Index differ')
    entry = [x for x in manifest['runtime_files'] if x['role'] == 'ENTRY_INDEX']
    require(len(entry) == 1 and entry[0]['sha256'] == profile['manifests']['runtime-index.json']['sha256']
            and entry[0]['file_id'] == index['runtime_index_file_id'] == manifest['runtime_entry_point']['file_id'],
            'NATIVE_INDEX_IDENTITY', 'Entry index identity differs')
    return {'components': ids, 'formal_package_version': row['package_version'], 'registry_entry': row,
            'product_qa': 'NOT_REISSUED', 'baseline_bytes': 'VERIFIED_ALL_SEVEN_PLUS_INDEX'}


def table_parity(checkout, profile):
    import openpyxl
    from openpyxl.utils.cell import range_boundaries, get_column_letter
    results = []
    for component in profile['components']:
        data = document(component['runtime'], checkout)
        require(set(data['tables']) == set(component['tables']), 'NATIVE_TABLE_SCOPE', component['component_id'])
        path = Path(checkout) / component['authoring']['path']
        calculated = openpyxl.load_workbook(path, read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
        cells = formula_count = 0
        try:
            require(set(calculated.sheetnames) == set(data['tables']), 'NATIVE_TABLE_SCOPE', 'All workbook sheets must be captured')
            for name, table in data['tables'].items():
                require(table['address'] == component['tables'][name], 'NATIVE_TABLE_RANGE', name)
                left, top, right, bottom = range_boundaries(table['address'])
                values = table['values']
                require(len(values) == bottom-top+1 and all(len(row) == right-left+1 for row in values),
                        'NATIVE_TABLE_SHAPE', name)
                for ri, row in enumerate(calculated[name][table['address']]):
                    for ci, cell in enumerate(row):
                        value = cell.value.isoformat() if hasattr(cell.value, 'isoformat') else cell.value
                        coordinate = get_column_letter(left+ci)+str(top+ri)
                        require(value == values[ri][ci], 'NATIVE_TABLE_PARITY', name + '!' + coordinate)
                        cells += 1
                for row in formulas[name]:
                    for cell in row:
                        require(cell.data_type != 'e', 'NATIVE_FORMULA_ERROR', name + '!' + cell.coordinate)
                        if cell.value is not None:
                            require(top <= cell.row <= bottom and left <= cell.column <= right,
                                    'NATIVE_UNCAPTURED_CELL', name + '!' + cell.coordinate)
                        if cell.data_type == 'f':
                            formula_count += 1
                            require(calculated[name][cell.coordinate].value is not None, 'NATIVE_UNCACHED_FORMULA', name + '!' + cell.coordinate)
            results.append({'component_id': component['component_id'], 'status': 'PASS', 'cells': cells,
                            'tables': len(data['tables']), 'formula_count': formula_count, 'mismatches': 0,
                            'formula_errors': 0, 'semantics': 'NOT_EVALUATED_BY_CELL_PARITY'})
        finally:
            calculated.close()
            formulas.close()
    return {'components': results, 'total_cells': sum(x['cells'] for x in results),
            'total_tables': sum(x['tables'] for x in results), 'product_qa': 'NOT_REISSUED'}


def evidence_inventory(checkout, profile):
    results = []
    for component in profile['components']:
        data = document(component['runtime'], checkout)
        rows = data['tables']['90_Evidence']['values']
        headers = profile['evidence_headers']
        header_index = next((i for i, row in enumerate(rows) if any(h in row for h in headers['id'])), None)
        require(header_index is not None, 'NATIVE_EVIDENCE_SCHEMA', component['component_id'])
        header = rows[header_index]
        columns = {k: next((header.index(h) for h in names if h in header), None) for k, names in headers.items()}
        require(all(v is not None for v in columns.values()), 'NATIVE_EVIDENCE_SCHEMA', 'Evidence ID/source/locator/version required')
        evidence = []
        for number, row in enumerate(rows[header_index+1:], start=header_index+2):
            if not row[columns['id']]:
                continue
            source = str(row[columns['source']] or '')
            # Both native URL cells and native Drive:ID cells are preserved.
            ids = re.findall(r'(?:Drive:|/d/)([A-Za-z0-9_-]+)', source)
            evidence.append({'id': row[columns['id']], 'row': number, 'source': source,
                             'drive_file_ids': ids, 'locator': row[columns['locator']],
                             'version': row[columns['version']], 'original_values': row,
                             'semantic_review': 'NOT_REISSUED'})
        ids = [x['id'] for x in evidence]
        require(evidence and len(ids) == len(set(ids)), 'NATIVE_EVIDENCE_IDENTITY', 'Evidence IDs absent or duplicated')
        results.append({'component_id': component['component_id'], 'evidence': evidence,
                        'accepted_decision': 'NOT_RETURNED; no separate accepted Decision artifact in acquired package',
                        'controlled_unresolved_rows': [
                            {'table': name, 'row': i+1, 'values': row}
                            for name, table in data['tables'].items() for i, row in enumerate(table['values'])
                            if any(isinstance(v, str) and any(t in v for t in ('EXTERNAL_CHECK', 'MANUAL_CHECK', 'HOLD', 'UNCONFIRMED', 'PENDING')) for v in row)]})
    return {'components': results, 'price_policy': 'NULL/HOLD retained; no 0 substitution',
            'qa_boundary': 'Evidence row inventory; existence is not manufacturer fact verification'}


def reference_closure(checkout, profile, registry_csv):
    import csv
    import io
    rows = list(csv.DictReader(io.StringIO(registry_csv)))
    inventory = evidence_inventory(checkout, profile)
    by_file = {fid: row for row in rows for fid in row.get('authoring_file_id', '').split(';')}
    results = []
    for component in inventory['components']:
        refs = []
        for e in component['evidence']:
            for fid in e['drive_file_ids']:
                row = by_file.get(fid)
                refs.append({'evidence_id': e['id'], 'file_id': fid,
                             'current_authoring_row': {k: row[k] for k in ('manufacturer','series','package_version','master_status')} if row else None,
                             'state': 'CURRENT_REGISTRY_ID_MATCH_ONLY' if row else 'NOT_MATCHED_TO_CURRENT_AUTHORING',
                             'bytes_and_semantic_binding': 'REVIEW_REQUIRED'})
        results.append({'component_id': component['component_id'], 'refs': refs,
                        'integration_state': 'BLOCKED_DEPENDENCY_PROJECTION_AND_PRODUCT_QA_REQUIRED'})
    return {'components': results, 'auto_mapping': False, 'all_integration_ready': False,
            'source_review': 'COMPLETE_OLD_NEW_OFFICIAL_BYTES_AND_ALL_SCOPE_REVIEW_REQUIRED'}


def stage(contract, gate, work):
    require(gate in STAGES, 'NATIVE_OPERATION_UNSUPPORTED', 'No mutation/adoption/release stage admitted')
    checkout = Path(contract['checkout'])
    profile = load_profile(checkout, contract['profile'])
    require(contract['product_id'] == profile['product_id']
            and contract['target_gate'] == profile['admitted_target_gate']
            and contract['scope'] == profile['scope'], 'NATIVE_SCOPE_MISMATCH', 'One exact audit scope required')
    registry = verified(contract['native_registry'], checkout).decode('utf-8')
    result = {'status': 'PASS', 'scope': profile['scope'], 'product_id': profile['product_id'],
              'formal_mutations': 0, 'runtime_ui_mutations': 0, 'external_operations': 0}
    if gate == 'BASELINE_RESOLUTION':
        result.update(baseline(checkout, profile, registry))
    elif gate == 'NATIVE_TABLE_PARITY':
        result.update(table_parity(checkout, profile))
    elif gate == 'EVIDENCE_INVENTORY':
        result.update(evidence_inventory(checkout, profile))
    elif gate == 'REFERENCE_CLOSURE':
        result.update(reference_closure(checkout, profile, registry))
    elif gate == 'REVIEW_PACKET':
        prior = {name: json.loads((Path(work)/(name+'.json')).read_text()) for name in STAGES[:4]}
        result.update(state='TECHNICAL_CONNECTION_REVIEW_ONLY', profile=contract['profile'], evidence=prior,
                      unsupported_operations=profile['unsupported_operations'],
                      product_qa='NOT_REISSUED', formal_adoption='NOT_PREPARED', release='NOT_PREPARED',
                      human_decision='NOT_REQUESTED_UNTIL_PRODUCT_CANDIDATE_IS_REVIEWABLE')
        result['payload_sha256'] = sha(json.dumps(result, ensure_ascii=False, sort_keys=True, separators=(',', ':')).encode())
    else:
        for name in STAGES[:-1]:
            path = Path(work)/(name+'.json')
            saved = Path(contract['storage_root'])/contract['work_id']/(name+'.json')
            require(path.is_file() and saved.is_file() and path.read_bytes() == saved.read_bytes(),
                    'NATIVE_SAVE_READBACK_REQUIRED', name)
        result.update(state='REPLUS_PACKAGE_AUDITED', storage='LOCAL_SAVED_AND_READ_BACK',
                      cloud_sync='NOT_VERIFIED', product_qa='NOT_REISSUED',
                      formal_adoption=False, app_integration_ready=False, release_completed=False,
                      policy=profile['terminal_policy'])
    return result


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--contract', required=True)
    p.add_argument('--stage', required=True)
    p.add_argument('--work', required=True)
    p.add_argument('--out', required=True)
    a = p.parse_args()
    try:
        result = stage(json.loads(Path(a.contract).read_text()), a.stage, a.work)
    except (Rejected, OSError, ValueError, KeyError, TypeError, ImportError) as error:
        result = {'status': 'FAIL', 'blocking_reason': {'code': getattr(error, 'code', 'NATIVE_PACKAGE_FAILURE'), 'reason': str(error)},
                  'external_operations': 0}
    Path(a.out).write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n')
    # Shared validator consumes the typed status; a FAIL can never be a gate PASS.


if __name__ == '__main__':
    main()
