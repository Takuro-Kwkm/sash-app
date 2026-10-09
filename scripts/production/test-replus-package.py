#!/usr/bin/env python3
"""Actual native baseline and adversarial scope tests; no fixture product certification."""
import copy
import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
s = importlib.util.spec_from_file_location('replus_package', ROOT/'scripts/production/replus-package.py')
native = importlib.util.module_from_spec(s)
s.loader.exec_module(native)
PROFILE = json.loads((ROOT/'contracts/production/replus/profile.v1.json').read_text())
REGISTRY = (ROOT/'contracts/production/replus/baseline/product-registry.observed.csv').read_text()


class NativePackage(unittest.TestCase):
    def test_actual_all_seven_formal_identity(self):
        self.assertEqual(len(native.baseline(ROOT, PROFILE, REGISTRY)['components']), 7)

    def test_actual_all_cells_and_formula_caches(self):
        result = native.table_parity(ROOT, PROFILE)
        self.assertEqual(result['total_cells'], 19910)
        self.assertEqual(result['total_tables'], 220)
        self.assertEqual(result['product_qa'], 'NOT_REISSUED')

    def test_native_header_variants_and_drive_id_cells(self):
        result = native.evidence_inventory(ROOT, PROFILE)
        self.assertEqual([len(c['evidence']) for c in result['components']], [22,13,12,12,15,18,16])
        self.assertEqual(result['components'][-1]['evidence'][1]['drive_file_ids'], ['1wH1nhmsftl6y8Mp7y6twtoXSi6rxqxzT'])

    def test_missing_official_dependency_proof_is_not_product_qa(self):
        result = native.reference_closure(ROOT, PROFILE, REGISTRY)
        self.assertFalse(result['all_integration_ready'])
        self.assertFalse(result['auto_mapping'])
        self.assertTrue(all(x['integration_state'].startswith('BLOCKED') for x in result['components']))

    def test_current_registry_version_drift_rejected(self):
        with self.assertRaisesRegex(native.Rejected, 'Registry/package'):
            native.baseline(ROOT, PROFILE, REGISTRY.replace('LIXIL,リプラス,v1.1,', 'LIXIL,リプラス,v999,'))

    def test_package_and_component_versions_are_distinct(self):
        changed = copy.deepcopy(PROFILE)
        changed['components'][0]['component_version'] = 'v1.1'
        with self.assertRaisesRegex(native.Rejected, 'Authoring/Runtime/Index'):
            native.baseline(ROOT, changed, REGISTRY)

    def test_missing_or_duplicate_component_rejected(self):
        changed = copy.deepcopy(PROFILE)
        changed['components'].pop()
        with self.assertRaises(native.Rejected):
            native.baseline(ROOT, changed, REGISTRY)

    def test_changed_workbook_hash_rejected_before_cell_qa(self):
        changed = copy.deepcopy(PROFILE)
        changed['components'][0]['authoring']['sha256'] = '0'*64
        with self.assertRaises(native.Rejected):
            native.baseline(ROOT, changed, REGISTRY)

    def test_out_of_scope_runtime_table_rejected(self):
        changed = copy.deepcopy(PROFILE)
        changed['components'][0]['tables'].pop('10B_現調項目')
        with self.assertRaises(native.Rejected):
            native.table_parity(ROOT, changed)

    def test_no_mutating_native_stages(self):
        for gate in ('FORMAL_CHANGE_ADOPTION','QA_READY','RELEASE','DEPLOYMENT','BUILD'):
            with self.subTest(gate=gate), self.assertRaises(native.Rejected):
                native.stage({}, gate, '.')

    def test_ci_admission_is_bounded_and_hash_bound(self):
        config = json.loads((ROOT/'contracts/production/work-connections.v2.json').read_text())
        entry = next(c for c in config['products'] if c['product_id'] == PROFILE['product_id'])
        validators = json.loads((ROOT/'contracts/production/validator-registry.v2.json').read_text())['products'][PROFILE['product_id']]
        admission = entry['existing_package_profile']
        self.assertEqual(admission['sha256'], native.sha((ROOT/admission['path']).read_bytes()))
        self.assertEqual(entry['default_gate'], PROFILE['admitted_target_gate'])
        self.assertEqual(validators['profile'], admission)
        self.assertEqual(validators['scope'], 'EXISTING_PACKAGE_AUDIT_ONLY')
        self.assertEqual(validators['mutation_policy'], 'NONE')
        self.assertFalse(entry['safe_failure']['fallback'])

    def test_quarantined_correction_preserves_all_unaffected_cells(self):
        import openpyxl
        folder = ROOT/'contracts/production/replus/candidates/cut-mall-r1-v2'
        plan = json.loads((folder/'patch-plan.json').read_text())
        allowed = {(p['sheet'],p['address']):p for p in plan['patches']}
        before = openpyxl.load_workbook(ROOT/'contracts/production/replus/baseline/RPL_CUT_MALL.xlsx', data_only=False)
        after = openpyxl.load_workbook(folder/'authoring-candidate.xlsx', data_only=False)
        try:
            self.assertEqual(before.sheetnames,after.sheetnames)
            changes = 0
            for sheet in before:
                target = after[sheet.title]
                for row in range(1,max(sheet.max_row,target.max_row)+1):
                    for col in range(1,max(sheet.max_column,target.max_column)+1):
                        a,b = sheet.cell(row,col),target.cell(row,col)
                        if a.value != b.value:
                            patch = allowed[(sheet.title,a.coordinate)]
                            self.assertEqual((a.value,b.value),(patch['before'],patch['after']))
                            changes += 1
                        if a.value is not None:
                            for key in ('font','fill','border','alignment','number_format','protection'):
                                self.assertEqual(str(getattr(a,key)),str(getattr(b,key)))
            self.assertEqual(changes,len(plan['patches']))
        finally:
            before.close();after.close()

    def test_quarantined_runtime_cannot_masquerade_as_ready(self):
        folder = ROOT/'contracts/production/replus/candidates/cut-mall-r1-v2'
        runtime = json.loads((folder/'runtime-candidate.json').read_text())
        manifest = json.loads((folder/'candidate-manifest.json').read_text())
        self.assertEqual(runtime['runtime_status'],'CANDIDATE_BLOCKED')
        self.assertFalse(runtime['automatic_orderability'])
        self.assertFalse(manifest['qa_ready'])
        self.assertFalse(manifest['formal_adopted'])
        self.assertEqual(len(manifest['carry_forward']),6)
        for ref in manifest['members']:
            data = (folder/ref['path']).read_bytes()
            self.assertEqual((len(data),native.sha(data)),(ref['bytes'],ref['sha256']))


if __name__ == '__main__':
    unittest.main()
