import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

ROOT=Path(__file__).resolve().parents[2]
spec=importlib.util.spec_from_file_location('native',ROOT/'scripts/production/sheet-package-change.py')
native=importlib.util.module_from_spec(spec);spec.loader.exec_module(native)


class NativeBoundaryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import os
        native.setup(os.environ['SHARED_HARNESS_ROOT'])

    def test_changed_bytes_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'source';path.write_bytes(b'old')
            ref={'path':str(path),'sha256':native.sha(b'old')}
            path.write_bytes(b'new')
            with self.assertRaises(native.Blocked):native.verified(ref)

    def test_symlinks_are_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'source';path.write_bytes(b'actual');link=Path(directory)/'link';link.symlink_to(path)
            with self.assertRaises(native.Blocked):native.verified({'path':str(link),'sha256':native.sha(b'actual')})

    def test_cell_changes_include_removals_formulas_and_added_sheets(self):
        old={'Product':{'A1':'value','A2':'=1+1','A3':'keep'}}
        new={'Product':{'A1':'value','A2':'=2+1'},'Evidence':{'A1':'source'}}
        self.assertEqual(native.delta(old,new),{'Evidence!A1':{'before':None,'after':'source'},'Product!A2':{'before':'=1+1','after':'=2+1'},'Product!A3':{'before':'keep','after':None}})

    def test_formula_errors_are_rejected(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'test.xlsx';book=openpyxl.Workbook();book.active['A1']='#REF!';book.save(path);book.close()
            with self.assertRaises(native.Blocked):native.workbook({'path':str(path),'sha256':native.sha(path.read_bytes())})

    def test_empty_cells_are_normal(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'test.xlsx';book=openpyxl.Workbook();book.active['C3']='value';book.save(path);book.close()
            self.assertEqual(native.workbook({'path':str(path),'sha256':native.sha(path.read_bytes())}),{'Sheet':{'C3':'value'}})

    def test_review_cannot_ignore_an_added_cell(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory)
            refs=[]
            for name,value in [('baseline','same'),('candidate','changed')]:
                book=openpyxl.Workbook();book.active['A1']=value;path=root/(name+'.xlsx');book.save(path);book.close();refs.append({'path':str(path),'sha256':native.sha(path.read_bytes())})
            review={'status':'COMPLETE_REVIEWED','review_ref':'TECHNICAL_UNIT_TEST_ONLY','product_id':'TEST','scope_id':'TEST','baseline_sha256':refs[0]['sha256'],'candidate_sha256':refs[1]['sha256'],
                    'sections':{'scope':'REVIEWED'},'unresolved_in_scope':[],'source_hashes':{},'cell_changes':{}}
            path=root/'review.json';path.write_text(json.dumps(review))
            c={'review':{'path':str(path),'sha256':native.sha(path.read_bytes())},'baseline_authoring':refs[0],'candidate_authoring':refs[1],'product_id':'TEST','scope_id':'TEST','sources':[]}
            with self.assertRaises(native.Blocked) as error:native.review(c,{'review_sections':['scope'],'native_sheets':['Sheet']})
            self.assertEqual(error.exception.data['code'],'UNREVIEWED_CELL_CHANGE')


if __name__=='__main__':unittest.main()
